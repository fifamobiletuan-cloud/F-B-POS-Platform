import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}>'
                      f'<w:top w:w="{top}" w:type="dxa"/>'
                      f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
                      f'<w:left w:w="{left}" w:type="dxa"/>'
                      f'<w:right w:w="{right}" w:type="dxa"/>'
                      f'</w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="CCCCCC"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'<w:tblBorders {nsdecls("w")}>'
                        f'<w:top w:val="single" w:sz="6" w:space="0" w:color="333333"/>'
                        f'<w:left w:val="none"/>'
                        f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="333333"/>'
                        f'<w:right w:val="none"/>'
                        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
                        f'<w:insideV w:val="none"/>'
                        f'</w:tblBorders>')
    tblPr.append(borders)

doc = docx.Document()

# Căn lề chuẩn văn bản học thuật (Nghị định 30 / TCVN: Trái 3cm, Phải 2cm, Trên 2cm, Dưới 2cm)
for section in doc.sections:
    section.top_margin = Inches(0.79)     # 2.0 cm
    section.bottom_margin = Inches(0.79)  # 2.0 cm
    section.left_margin = Inches(1.18)    # 3.0 cm
    section.right_margin = Inches(0.79)   # 2.0 cm

# Thiết lập Style chuẩn Times New Roman
style_normal = doc.styles['Normal']
style_normal.font.name = 'Times New Roman'
style_normal.font.size = Pt(13)
style_normal.font.color.rgb = RGBColor(0, 0, 0)
style_normal.paragraph_format.line_spacing = 1.3
style_normal.paragraph_format.space_after = Pt(6)
style_normal.paragraph_format.space_before = Pt(0)
style_normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

def add_title_1(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    run.font.bold = True
    return p

def add_title_2(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(13)
    run.font.bold = True
    return p

def add_title_3(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.italic = True
    return p

def add_para(text, indent=True):
    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.3
    if indent:
        p.paragraph_format.first_line_indent = Inches(0.39) # Thụt đầu dòng 1cm chuẩn sinh viên viết
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(13)
    return p

# ==================== TRANG BÌA ====================
p_truong = doc.add_paragraph()
p_truong.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_truong.paragraph_format.space_after = Pt(2)
r_tr = p_truong.add_run("BỘ GIÁO DỤC VÀ ĐÀO TẠO\nTRƯỜNG ĐẠI HỌC ........................................\nKHOA CÔNG NGHỆ THÔNG TIN")
r_tr.font.name = 'Times New Roman'
r_tr.font.size = Pt(12.5)
r_tr.font.bold = True

p_sao = doc.add_paragraph()
p_sao.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_sao.paragraph_format.space_after = Pt(40)
r_s = p_sao.add_run("-------------------***-------------------")
r_s.font.bold = True

p_mon = doc.add_paragraph()
p_mon.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_mon.paragraph_format.space_after = Pt(8)
r_mon = p_mon.add_run("ĐỀ CƯƠNG BÁO CÁO ĐỀ TÀI\nCHUYÊN ĐỀ CÔNG NGHỆ THÔNG TIN")
r_mon.font.name = 'Times New Roman'
r_mon.font.size = Pt(17)
r_mon.font.bold = True

p_huong = doc.add_paragraph()
p_huong.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_huong.paragraph_format.space_after = Pt(36)
r_huong = p_huong.add_run("(Định hướng tiếp tục phát triển làm Đồ án Tốt nghiệp)")
r_huong.font.name = 'Times New Roman'
r_huong.font.size = Pt(12)
r_huong.font.italic = True

p_ten_nhan = doc.add_paragraph()
p_ten_nhan.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_ten_nhan.paragraph_format.space_after = Pt(4)
r_tn = p_ten_nhan.add_run("TÊN ĐỀ TÀI:")
r_tn.font.name = 'Times New Roman'
r_tn.font.size = Pt(13)
r_tn.font.bold = True

p_ten = doc.add_paragraph()
p_ten.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_ten.paragraph_format.space_after = Pt(8)
r_ten = p_ten.add_run("ỨNG DỤNG GỌI MÓN TẠI BÀN BẰNG MÃ QR VÀ HỆ THỐNG QUẢN LÝ VẬN HÀNH QUÁN TRÀ SỮA DỰA TRÊN FLUTTER VÀ SPRING BOOT")
r_ten.font.name = 'Times New Roman'
r_ten.font.size = Pt(15)
r_ten.font.bold = True

p_tienganh = doc.add_paragraph()
p_tienganh.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_tienganh.paragraph_format.space_after = Pt(70)
r_ta = p_tienganh.add_run("(Smart QR Ordering and Milk Tea Store Management System)")
r_ta.font.name = 'Times New Roman'
r_ta.font.size = Pt(12)
r_ta.font.italic = True

# Khung thông tin sinh viên
table_sv = doc.add_table(rows=4, cols=2)
table_sv.alignment = WD_TABLE_ALIGNMENT.CENTER
table_sv.autofit = False

sv_info = [
    ("Giảng viên hướng dẫn:", "ThS/TS. ...................................................."),
    ("Sinh viên thực hiện:", "................................................................"),
    ("Mã số sinh viên:", "................................................................"),
    ("Ngành học / Lớp:", "Kỹ thuật phần mềm / ................................")
]

for row_idx, (label, val) in enumerate(sv_info):
    row = table_sv.rows[row_idx]
    c_l = row.cells[0]
    c_r = row.cells[1]
    c_l.width = Inches(2.2)
    c_r.width = Inches(3.8)
    
    pl = c_l.paragraphs[0]
    pl.paragraph_format.space_after = Pt(4)
    rl = pl.add_run(label)
    rl.font.name = 'Times New Roman'
    rl.font.size = Pt(12.5)
    rl.font.bold = True
    
    pr = c_r.paragraphs[0]
    pr.paragraph_format.space_after = Pt(4)
    rr = pr.add_run(val)
    rr.font.name = 'Times New Roman'
    rr.font.size = Pt(12.5)

p_nam = doc.add_paragraph()
p_nam.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_nam.paragraph_format.space_before = Pt(70)
r_nam = p_nam.add_run("Năm học 2025 - 2026")
r_nam.font.name = 'Times New Roman'
r_nam.font.size = Pt(12)

doc.add_page_break()

# ==================== NỘI DUNG TỰ NHIÊN, KHÔNG VĂN MẪU AI ====================

add_title_1("PHẦN 1. LÝ DO CHỌN ĐỀ TÀI VÀ MỤC TIÊU THỰC HIỆN")

add_title_2("1.1. Xuất phát điểm từ thực tế đời sống")
add_para("Hiện nay, mô hình kinh doanh quán trà sữa, trà trái cây và các món ăn vặt xung quanh các trường đại học, khu ký túc xá phát triển rất nhanh để phục vụ nhu cầu học tập, gặp gỡ của sinh viên. Xuất phát từ chính trải nghiệm của bản thân khi đi uống nước và quan sát tại các quán quen, em nhận thấy quy trình phục vụ tại hầu hết các quán quy mô vừa và nhỏ hiện nay vẫn còn gặp nhiều bất cập.")

add_para("Thứ nhất là tình trạng quá tải vào khung giờ cao điểm, đặc biệt là khoảng từ 19h đến 21h hàng ngày hoặc các ngày cuối tuần. Khách vào quán thường phải đứng xếp hàng chờ đợi khá lâu tại quầy để gọi món. Nhiều bạn vào bàn ngồi trước thì nhân viên bận rộn không kịp mang thực đơn ra ghi món, dẫn đến việc khách phải gọi với hoặc đợi lâu gây khó chịu.")

add_para("Thứ hai, đặc thù của món trà sữa khác hẳn so với cà phê hay nước ngọt đóng chai ở chỗ có rất nhiều tùy chỉnh riêng cho từng người: chọn size M hay size L, chọn mức đường (0%, 30%, 50%, 70%, 100%), chọn mức đá (không đá, ít đá, đá bình thường) và rất nhiều loại topping khác nhau như trân châu đen, thạch củ năng, pudding trứng, thạch phô mai. Việc nhân viên ghi chép tay trên giấy nhớ hoặc nhớ miệng khi quán đông khách rất dễ dẫn đến sai sót: nhầm size, quên topping hoặc cho sai lượng đường. Khi làm sai, quán buộc phải làm lại ly khác, gây lãng phí nguyên liệu và mất thời gian của khách.")

add_para("Thứ ba là vấn đề quản lý nội bộ của quán. Phần lớn các chủ quán nhỏ hiện nay chỉ theo dõi số tiền thu vào hàng ngày, còn việc chi tiêu lặt vặt như mua đá lạnh, mua thêm trà, sữa đặc, ly nhựa, ống hút hay tiền sửa máy dập nắp thì thường chi tiền mặt ra mà không có chỗ ghi chép tập trung. Đến cuối tháng, chủ quán rất khó biết được lợi nhuận thực tế sau khi trừ đi các chi phí phát sinh là bao nhiêu. Bên cạnh đó, nhân viên quán chủ yếu là sinh viên làm theo ca (Part-time), thường xuyên phải đổi ca học và xin đổi ca làm việc cho nhau, việc theo dõi giờ công qua nhóm chat hay sổ tay rất dễ nhầm lẫn khi tính lương.")

add_title_2("1.2. Giải pháp đề xuất của sinh viên")
add_para("Để giải quyết trực tiếp các vấn đề thực tế trên, em đề xuất xây dựng một hệ thống gồm hai ứng dụng cụ thể:")
add_para("Một là ứng dụng web gọi món dành cho khách hàng: Khách vào quán ngồi tại bàn nào thì mở camera điện thoại quét mã QR dán trên bàn đó. Trình duyệt trên điện thoại sẽ tự động mở trang thực đơn đúng số bàn của khách mà không cần cài đặt thêm bất kỳ ứng dụng nào vào máy. Khách có thể tự do xem ảnh món, chọn lượng đường, đá, chọn topping theo ý thích và bấm gửi đơn hàng trực tiếp về quầy pha chế, hoặc chọn quét mã VietQR để thanh toán luôn.", indent=True)
add_para("Hai là ứng dụng di động dành riêng cho cửa hàng: Được cài đặt trên điện thoại hoặc máy tính bảng của nhân viên và chủ quán. Ứng dụng này phục vụ việc nhận đơn từ bàn khách gửi về theo thời gian thực để nhân viên pha chế đúng yêu cầu; quản lý sơ đồ bàn xem bàn nào đang có khách; ghi lại các khoản chi tiêu mua sắm hàng ngày; xem thống kê doanh thu - chi phí - lợi nhuận và theo dõi lịch xếp ca cũng như chấm công của nhân viên.", indent=True)

add_title_2("1.3. Mục tiêu và phạm vi của đề tài")
add_para("Mục tiêu của em khi thực hiện đề tài này gồm hai khía cạnh:")
add_para("Về mặt ứng dụng thực tế: Xây dựng một công cụ nhỏ gọn, dễ dùng, phục vụ trực tiếp cho các quán trà sữa và đồ ăn vặt; giúp giảm tải việc đứng quầy ghi món của nhân viên, hạn chế tối đa việc làm nhầm món và giúp chủ quán quản lý rõ ràng cả tiền thu lẫn tiền chi hàng ngày.", indent=True)
add_para("Về mặt học tập và kỹ năng: Áp dụng các kiến thức đã học trong chương trình Công nghệ thông tin để tự tay xây dựng một hệ thống hoàn chỉnh từ Backend (Java Spring Boot, PostgreSQL) đến Frontend (Flutter); làm quen với quy trình kiểm thử API bài bản bằng Postman và quản lý mã nguồn bằng Git.", indent=True)

add_para("Phạm vi đề tài tập trung vào nghiệp vụ của một quán trà sữa - đồ ăn vặt cụ thể. Đối tượng sử dụng bao gồm khách hàng tại bàn, nhân viên pha chế/thu ngân và chủ quán.")

# ==================== PHẦN 2 ====================
add_title_1("PHẦN 2. KIẾN TRÚC KỸ THUẬT VÀ CÔNG NGHỆ LỰA CHỌN")

add_title_2("2.1. Ngăn xếp công nghệ sử dụng")
add_para("Dựa trên định hướng nghề nghiệp sau khi ra trường theo mảng Phát triển ứng dụng di động (Mobile App) và Lập trình Backend Java, em lựa chọn kiến trúc công nghệ như sau:")

add_para("1. Giao diện người dùng (Frontend): Sử dụng Flutter với ngôn ngữ Dart. Điểm mạnh lớn nhất mà em muốn khai thác ở Flutter là khả năng viết một lần nhưng chạy được trên nhiều nền tảng (Single Codebase). Cụ thể, em sử dụng Flutter Web để làm trang gọi món cho khách (khách quét QR mở ngay trên trình duyệt mà không phải tải app nặng máy) và Flutter Mobile để đóng gói thành ứng dụng Android cài lên điện thoại của nhân viên và chủ quán.", indent=True)

add_para("2. Phía máy chủ (Backend): Sử dụng Java 17 và nền tảng Spring Boot 3.x để viết các RESTful API. Spring Boot là framework chuẩn mực trong các doanh nghiệp hiện nay với cấu trúc phân tầng rõ ràng, khả năng xử lý đồng thời tốt và hệ sinh thái thư viện rất phong phú.", indent=True)

add_para("3. Cơ sở dữ liệu và Tầng ORM: Sử dụng PostgreSQL kết hợp Spring Data JPA và Hibernate. Vì hệ thống liên quan đến việc tính tiền hóa đơn, ghi nhận chi phí và chấm công, em chọn PostgreSQL để đảm bảo tính toàn vẹn dữ liệu (tuân thủ chuẩn ACID). Tầng Hibernate giúp ánh xạ trực tiếp các bảng dữ liệu thành các Class Java (Entity), giúp thao tác dữ liệu an toàn và hạn chế tối đa việc phải ghép chuỗi SQL thủ công.", indent=True)

add_para("4. Bảo mật và Phân quyền: Sử dụng Spring Security kết hợp JSON Web Token (JWT) và OAuth2. Hệ thống áp dụng cơ chế Stateless Session (không lưu phiên đăng nhập trên server). Phân quyền người dùng được chia thành các nhóm rõ ràng: khách hàng quét QR chỉ được phép xem thực đơn và gửi đơn; nhân viên (ROLE_STAFF) được xem màn hình pha chế và chấm công; chủ quán (ROLE_ADMIN) mới có quyền xem doanh thu, ghi chi tiêu và quản lý thực đơn.", indent=True)

add_para("5. Công cụ phát triển và kiểm thử: Sử dụng Postman để kiểm tra tính đúng đắn của từng API (mã phản hồi, dữ liệu trả về) trước khi ghép nối vào ứng dụng; sử dụng Git và GitHub để quản lý phiên bản mã nguồn theo từng nhánh tính năng.", indent=True)

add_title_2("2.2. Bảng tổng hợp công nghệ")

table_tech = doc.add_table(rows=8, cols=3)
table_tech.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(table_tech)

tech_rows = [
    ("Thành phần", "Công nghệ sử dụng", "Mục đích sử dụng"),
    ("Ứng dụng Khách hàng", "Flutter Web (Ngôn ngữ Dart)", "Trang web mở nhanh qua mã QR để khách chọn món tại bàn"),
    ("Ứng dụng Cửa hàng", "Flutter Mobile (Android/iOS)", "App cài trên điện thoại nhân viên và chủ quán để quản lý"),
    ("Backend REST API", "Java 17, Spring Boot 3.x", "Xử lý logic nghiệp vụ, tính tiền, phân quyền và cung cấp API"),
    ("Cơ sở dữ liệu", "PostgreSQL 16", "Lưu trữ dữ liệu bàn ăn, thực đơn, đơn hàng, thu chi và ca làm"),
    ("Tầng ORM", "Spring Data JPA / Hibernate", "Ánh xạ cơ sở dữ liệu sang đối tượng Java, quản lý giao dịch"),
    ("Bảo mật", "Spring Security, JWT", "Xác thực tài khoản và phân quyền truy cập theo vai trò"),
    ("Kiểm thử & Mã nguồn", "Postman, Git / GitHub", "Kiểm thử API tự động và quản lý lịch sử viết mã nguồn")
]

for row_idx, data in enumerate(tech_rows):
    row = table_tech.rows[row_idx]
    for col_idx, text in enumerate(data):
        c = row.cells[col_idx]
        set_cell_margins(c, top=100, bottom=100, left=120, right=120)
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.space_before = Pt(2)
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(11.5)
        if row_idx == 0:
            set_cell_background(c, "EAEAEA")
            r.font.bold = True
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            if col_idx == 0:
                r.font.bold = True

# ==================== PHẦN 3 ====================
add_title_1("PHẦN 3. ĐẶC TẢ CÁC CHỨC NĂNG CỦA HỆ THỐNG")

add_title_2("3.1. Phân hệ Web dành cho Khách hàng (Quét QR tại bàn)")
add_para("Khách hàng vào quán chỉ sử dụng điện thoại thông minh quét mã QR dán trên bàn, với các chức năng cụ thể:")
add_para("1. Nhận diện vị trí bàn: Khách quét mã QR có dạng liên kết kèm theo mã nhận diện bàn (ví dụ: .../?table=B05). Hệ thống tự động xác định khách đang ngồi bàn số 5 và gắn mã bàn này vào đơn hàng.", indent=True)
add_para("2. Xem thực đơn và chọn món: Thực đơn hiển thị hình ảnh, tên món, giá tiền và phân loại rõ ràng (Trà sữa, Trà trái cây, Ăn vặt). Khách có thể tìm kiếm món theo tên.", indent=True)
add_para("3. Tùy chọn chi tiết từng ly (Customization): Khi bấm vào một ly trà sữa, giao diện sẽ cho phép khách chọn kích cỡ (Size M mặc định, Size L +6.000đ); chọn mức độ ngọt theo phần trăm (0%, 30%, 50%, 70%, 100%); chọn lượng đá (nóng, không đá, ít đá, bình thường) và chọn thêm các loại topping như trân châu đen, thạch phô mai, pudding. Tiền topping sẽ tự động cộng dồn vào giá của ly đó.", indent=True)
add_para("4. Giỏ hàng và Đặt món: Khách kiểm tra lại danh sách các ly đã chọn cùng các ghi chú kèm theo. Sau khi xác nhận, đơn hàng được gửi thẳng về hệ thống của quán.", indent=True)
add_para("5. Thanh toán: Khách có thể chọn quét mã VietQR để chuyển khoản thanh toán với nội dung và số tiền đã được tạo sẵn tự động, hoặc chọn thanh toán tiền mặt tại quầy khi ra về.", indent=True)
add_para("6. Theo dõi trạng thái: Khách có thể nhìn thấy đồ uống của mình đang ở bước nào: Đã nhận đơn -> Đang pha chế -> Đã hoàn thành.", indent=True)

add_title_2("3.2. Phân hệ Ứng dụng Di động dành cho Cửa hàng")
add_para("Ứng dụng dành cho nhân viên và chủ quán được tổ chức thành các màn hình chức năng rõ ràng:")

add_title_3("a) Màn hình quầy pha chế (Bếp)")
add_para("Khi có bàn gửi đơn, màn hình của nhân viên pha chế sẽ lập tức hiện đơn mới theo đúng số bàn, kèm chi tiết từng ly (ví dụ: Bàn 02 - 1 Trà sữa Ô long size L, 50% đường, thêm trân châu trắng). Nhân viên pha xong ly nào hoặc bàn nào thì chạm vào nút để chuyển trạng thái sang Đã hoàn thành.")

add_title_3("b) Màn hình quản lý sơ đồ bàn")
add_para("Giao diện trực quan hiển thị danh sách các bàn trong quán kèm màu sắc nhận biết: bàn màu xanh là bàn trống, bàn màu đỏ là bàn đang có khách ngồi và đang dùng đồ. Nhân viên có thể hỗ trợ khách đổi bàn hoặc tách gộp bàn khi cần thiết.")

add_title_3("c) Quản lý chi tiêu hàng ngày của quán")
add_para("Cho phép người quản lý hoặc nhân viên được ủy quyền nhập nhanh các khoản tiền chi ra trong ngày: mua nguyên liệu (sữa, trà, bột béo, trân châu), mua vật dụng (ly nhựa, màng ép, ống hút), tiền điện nước hoặc tiền đá lạnh mua lẻ ngoài đại lý. Ứng dụng hỗ trợ chụp lại ảnh hóa đơn mua hàng để lưu trữ làm bằng chứng đối soát.")

add_title_3("d) Báo cáo doanh số và lợi nhuận")
add_para("Hệ thống tự động tổng hợp số tiền bán được trong ngày, tuần, tháng; đồng thời trừ đi các khoản chi phí đã nhập ở trên để đưa ra con số lợi nhuận thực tế. Chủ quán cũng xem được danh sách các món bán chạy nhất để biết xu hướng khách thích uống món gì.")

add_title_3("e) Quản lý ca làm việc và chấm công nhân viên")
add_para("Chủ quán tạo trước các ca làm việc chuẩn của quán (Ca sáng: 7h30 - 12h30, Ca chiều: 12h30 - 17h30, Ca tối: 17h30 - 22h30) và phân công lịch làm cho nhân viên theo từng tuần. Nhân viên khi đến quán làm việc chỉ cần mở app bấm 'Vào ca' và bấm 'Hết ca' khi tan làm. Cuối tháng, hệ thống tự cộng tổng số giờ làm việc thực tế của từng bạn để tính lương chính xác.")

# ==================== PHẦN 4 ====================
add_title_1("PHẦN 4. THIẾT KẾ DỮ LIỆU VÀ GIAO TIẾP HỆ THỐNG")

add_title_2("4.1. Thiết kế Cơ sở dữ liệu (PostgreSQL)")
add_para("Dựa trên các yêu cầu nghiệp vụ trên, em thiết kế cơ sở dữ liệu gồm 12 bảng chính, được liên kết chặt chẽ qua khóa ngoại và ràng buộc toàn vẹn dữ liệu:")

table_db = doc.add_table(rows=13, cols=3)
table_db.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(table_db)

db_rows = [
    ("Bảng dữ liệu", "Các cột chính", "Mô tả nội dung lưu trữ"),
    ("users", "id, username, password, full_name, phone, role_id", "Thông tin tài khoản của nhân viên và chủ quán"),
    ("roles", "id, name (ROLE_ADMIN, ROLE_STAFF)", "Bảng định nghĩa quyền hạn truy cập"),
    ("dining_tables", "id, table_number, qr_token, status", "Danh sách các bàn trong quán và mã token của mã QR"),
    ("categories", "id, name, display_order", "Phân loại món: Trà sữa, Trà hoa quả, Đồ ăn vặt"),
    ("products", "id, category_id, name, base_price, image_url, is_active", "Danh mục các món ăn và đồ uống của quán"),
    ("product_options", "id, name, option_type, additional_price", "Các lựa chọn đi kèm: Size, Đường, Đá, Topping"),
    ("orders", "id, table_id, total_amount, payment_method, status, created_at", "Thông tin đơn hàng, số tiền và trạng thái xử lý"),
    ("order_items", "id, order_id, product_id, quantity, unit_price", "Danh sách các món được gọi trong một đơn hàng"),
    ("order_item_options", "id, order_item_id, product_option_id, price", "Các tùy chọn đường, đá, topping đi kèm từng món"),
    ("expenses", "id, title, amount, category, invoice_image, created_by", "Các phiếu ghi chép chi tiêu tiền mặt hàng ngày"),
    ("shifts", "id, shift_name, start_time, end_time, hourly_rate", "Danh mục ca làm việc mẫu và mức lương mỗi giờ"),
    ("work_schedules", "id, user_id, shift_id, work_date, check_in, check_out", "Lịch xếp ca và dữ liệu chấm công của nhân viên")
]

for row_idx, data in enumerate(db_rows):
    row = table_db.rows[row_idx]
    for col_idx, text in enumerate(data):
        c = row.cells[col_idx]
        set_cell_margins(c, top=90, bottom=90, left=120, right=120)
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.space_before = Pt(2)
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(11)
        if row_idx == 0:
            set_cell_background(c, "EAEAEA")
            r.font.bold = True
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            if col_idx == 0:
                r.font.bold = True

add_title_2("4.2. Danh sách các RESTful API chính")
add_para("Backend Spring Boot cung cấp các điểm cuối API theo chuẩn REST, dữ liệu truyền nhận ở định dạng JSON:")

table_api = doc.add_table(rows=10, cols=4)
table_api.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(table_api)

api_rows = [
    ("Phương thức", "Đường dẫn API (Endpoint)", "Chức năng thực hiện", "Quyền truy cập"),
    ("POST", "/api/v1/auth/login", "Đăng nhập tài khoản, trả về mã xác thực JWT", "Mọi người"),
    ("GET", "/api/v1/customer/tables/{token}", "Đọc mã QR, trả về số bàn tương ứng", "Khách tại bàn"),
    ("GET", "/api/v1/customer/menu", "Lấy toàn bộ thực đơn kèm danh sách topping", "Khách tại bàn"),
    ("POST", "/api/v1/customer/orders", "Gửi đơn gọi món của bàn về hệ thống", "Khách tại bàn"),
    ("GET", "/api/v1/staff/orders", "Lấy danh sách các đơn hàng cần pha chế", "Nhân viên, Quản lý"),
    ("PATCH", "/api/v1/staff/orders/{id}", "Đổi trạng thái món sang đang làm hoặc đã xong", "Nhân viên, Quản lý"),
    ("POST", "/api/v1/expenses", "Tạo phiếu ghi chép một khoản chi tiêu mới", "Nhân viên, Quản lý"),
    ("POST", "/api/v1/shifts/check-in", "Nhân viên bấm chấm công vào ca làm việc", "Nhân viên"),
    ("GET", "/api/v1/reports/profit", "Lấy báo cáo doanh thu, chi phí và lợi nhuận", "Chỉ Quản lý (Admin)")
]

for row_idx, data in enumerate(api_rows):
    row = table_api.rows[row_idx]
    for col_idx, text in enumerate(data):
        c = row.cells[col_idx]
        set_cell_margins(c, top=90, bottom=90, left=120, right=120)
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.space_before = Pt(2)
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(11)
        if row_idx == 0:
            set_cell_background(c, "EAEAEA")
            r.font.bold = True
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            if col_idx == 0:
                r.font.bold = True
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER

# ==================== PHẦN 5 ====================
add_title_1("PHẦN 5. KẾ HOẠCH THỰC HIỆN VÀ TÀI LIỆU THAM KHẢO")

add_title_2("5.1. Kế hoạch thực hiện trong môn Chuyên đề CNTT")
add_para("Trong phạm vi môn Chuyên đề CNTT của học kỳ này, em xác định mục tiêu hoàn thành sản phẩm cơ bản có thể chạy thực tế được (MVP), bao gồm:")
add_para("Giai đoạn 1 (3 tuần đầu): Hoàn thành đề cương chi tiết, khảo sát lại thực tế quy trình tại quán trà sữa mẫu; thiết kế chi tiết các bảng trong CSDL PostgreSQL và tạo các Entity Java tương ứng bằng Spring Data JPA.", indent=True)
add_para("Giai đoạn 2 (3 tuần tiếp theo): Viết xong toàn bộ các API Backend bằng Spring Boot cho phần bàn ăn, thực đơn, tạo đơn hàng, ghi chi tiêu và chấm công. Cài đặt bảo mật Spring Security và JWT. Kiểm tra kỹ lưỡng các API trên Postman để bảo đảm không có lỗi logic.", indent=True)
add_para("Giai đoạn 3 (3 tuần cuối): Xây dựng giao diện Flutter Web cho khách quét QR đặt món và giao diện Flutter App cho quán nhận đơn và chấm công; sau đó tiến hành ghép nối API, kiểm thử và viết báo cáo tổng kết môn học.", indent=True)

add_title_2("5.2. Hướng mở rộng phát triển thành Đồ án Tốt nghiệp")
add_para("Sau khi hoàn thành môn Chuyên đề CNTT, bước sang kỳ làm Đồ án Tốt nghiệp, em sẽ tiếp tục đào sâu kỹ thuật và hoàn thiện hệ thống với các tính năng nâng cao sau:")
add_para("Thứ nhất, tích hợp WebSocket: Khi khách tại bàn bấm gửi đơn hàng, màn hình quầy pha chế sẽ tự động rung chuông và hiển thị món mới ngay lập tức mà nhân viên không cần phải bấm làm mới (F5) ứng dụng.", indent=True)
add_para("Thứ hai, kết nối thanh toán tự động: Tích hợp dịch vụ ngân hàng để khi khách chuyển khoản qua mã VietQR, hệ thống sẽ tự động nhận diện tiền đã vào tài khoản và tự động xác nhận đơn hàng thành công.", indent=True)
add_para("Thứ ba, in hóa đơn trực tiếp: Kết nối ứng dụng Flutter với máy in nhiệt tại quầy qua mạng LAN hoặc Bluetooth để in phiếu gọi món cho quầy pha chế và in hóa đơn thanh toán cho khách.", indent=True)
add_para("Thứ tư, đóng gói và triển khai: Tìm hiểu cách đóng gói hệ thống bằng Docker và đưa ứng dụng lên máy chủ đám mây (Cloud Server) để chạy thử nghiệm thực tế.", indent=True)

add_title_2("5.3. Tài liệu tham khảo chính")
add_para("1. Tài liệu hướng dẫn chính thức của Spring Boot 3: https://docs.spring.io/spring-boot/docs/current/reference/html/", indent=True)
add_para("2. Tài liệu phát triển ứng dụng đa nền tảng Flutter & Dart: https://docs.flutter.dev/", indent=True)
add_para("3. Hướng dẫn quản trị và thiết kế cơ sở dữ liệu PostgreSQL 16: https://www.postgresql.org/docs/16/", indent=True)
add_para("4. Sách 'Spring in Action' (Tác giả Craig Walls, Nhà xuất bản Manning Publications).", indent=True)
add_para("5. Giáo trình môn Lập trình mạng và Phát triển ứng dụng trên thiết bị di động của trường.", indent=True)

try:
    output_file = r"d:\Doantotnghiep\Bao_Cao_De_Xuat_De_Tai_Chuyen_De.docx"
    doc.save(output_file)
    print(f"File updated successfully at: {output_file}")
except PermissionError:
    output_file = r"d:\Doantotnghiep\Bao_Cao_De_Xuat_De_Tai_Chuyen_De_Chuan.docx"
    doc.save(output_file)
    print(f"File updated successfully at: {output_file} (Original was locked)")
